"""
Извлечение текста из вложений (pptx, pdf, docx, xlsx) для передачи в GPT-анализ.
"""

from __future__ import annotations

import io

from pptx import Presentation

from logger import get_logger

logger = get_logger(__name__)

# Максимум символов из вложений, чтобы не раздувать промпт
MAX_ATTACHMENT_TEXT = 6000


def _parse_pptx(data: bytes) -> str:
    """Извлекает весь текст из PowerPoint-презентации."""
    prs = Presentation(io.BytesIO(data))
    texts: list[str] = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    line = paragraph.text.strip()
                    if line:
                        slide_texts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells if cell.text.strip()
                    )
                    if row_text:
                        slide_texts.append(row_text)
        if slide_texts:
            texts.append(f"[Слайд {slide_num}]\n" + "\n".join(slide_texts))
    return "\n\n".join(texts)


def _parse_pdf(data: bytes) -> str:
    """Извлекает текст из PDF."""
    import fitz  # pymupdf

    doc = fitz.open(stream=data, filetype="pdf")
    texts: list[str] = []
    for page_num, page in enumerate(doc, start=1):
        text = page.get_text().strip()
        if text:
            texts.append(f"[Страница {page_num}]\n{text}")
    doc.close()
    return "\n\n".join(texts)


def _parse_docx(data: bytes) -> str:
    """Извлекает текст из Word-документа."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    texts: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            texts.append(text)
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                texts.append(row_text)
    return "\n".join(texts)


def _parse_xlsx(data: bytes) -> str:
    """Извлекает текст из Excel-файла."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    texts: list[str] = []
    for sheet in wb.worksheets:
        sheet_lines: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                sheet_lines.append(" | ".join(cells))
        if sheet_lines:
            texts.append(f"[Лист: {sheet.title}]\n" + "\n".join(sheet_lines))
    wb.close()
    return "\n\n".join(texts)


# Маппинг расширение → парсер
_PARSERS: dict[str, callable] = {
    ".pptx": _parse_pptx,
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".xlsx": _parse_xlsx,
}


def parse_attachment(filename: str, data: bytes) -> str | None:
    """
    Пытается извлечь текст из вложения по расширению.
    Возвращает текст или None, если формат не поддерживается.
    """
    ext = ""
    dot_pos = filename.rfind(".")
    if dot_pos != -1:
        ext = filename[dot_pos:].lower()

    parser = _PARSERS.get(ext)
    if parser is None:
        return None

    try:
        text = parser(data)
        return text.strip() or None
    except Exception as e:
        logger.warning(
            "Failed to parse attachment {fn}: {err}",
            fn=filename, err=e,
        )
        return None


def extract_attachments_text(raw_attachments: list[dict]) -> str:
    """
    Парсит все поддерживаемые вложения и возвращает объединённый текст.
    Обрезает до MAX_ATTACHMENT_TEXT символов.
    """
    parts: list[str] = []
    for att in raw_attachments:
        text = parse_attachment(att["filename"], att["data"])
        if text:
            parts.append(f"--- {att['filename']} ---\n{text}")
            logger.info(
                "Parsed attachment {fn}: {chars} chars",
                fn=att["filename"], chars=len(text),
            )

    if not parts:
        return ""

    combined = "\n\n".join(parts)
    if len(combined) > MAX_ATTACHMENT_TEXT:
        combined = combined[:MAX_ATTACHMENT_TEXT] + "\n[...текст обрезан]"
    return combined
