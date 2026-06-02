import io
import uuid
from datetime import timedelta

from openpyxl import load_workbook

from models.incoming_emails import IncomingEmail, EmailStatusEnum
from models.announcement import get_astana_time

HDR = {"X-User-Id": str(uuid.uuid4())}


async def test_requires_auth(client):
    r = await client.get("/announcements/reports/messages.xlsx",
                         params={"date_from": "2026-06-01", "date_to": "2026-06-30"})
    assert r.status_code == 401


async def test_bad_period_returns_400(client):
    r = await client.get("/announcements/reports/messages.xlsx",
                         params={"date_from": "2026-06-30", "date_to": "2026-06-01"}, headers=HDR)
    assert r.status_code == 400


async def test_returns_xlsx(client, session):
    e = IncomingEmail(outlook_id=f"o-{uuid.uuid4()}", subject="s", body="b",
                      sender_email="sd_info@Fortebank.com",
                      received_at=get_astana_time(), status=EmailStatusEnum.DONE)
    session.add(e); await session.flush()
    r = await client.get(
        "/announcements/reports/messages.xlsx",
        params={"date_from": (get_astana_time() - timedelta(days=1)).date().isoformat(),
                "date_to": (get_astana_time() + timedelta(days=1)).date().isoformat()},
        headers=HDR,
    )
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Сводка", "Сообщения"]
