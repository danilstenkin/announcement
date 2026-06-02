import io
from datetime import datetime
from uuid import uuid4

from openpyxl import load_workbook

from services.reporting import MessageRow, Summary
from services.report_export import build_report_xlsx


def _row():
    return MessageRow(
        email_id=uuid4(), received_at=datetime(2026, 6, 1, 10, 0),
        sender_email="sd_info@fortebank.com", source="ServiceDesk", subject="Тема",
        email_status="DONE", processing_type="AUTO", ticket_status=None,
        assignee_name=None, publish_at=None, addressed=True, idle_hours=None,
        announcement_id=None, announcement_title=None, published_at=None,
        cc_scope="для КЦ",
    )


def test_xlsx_has_two_sheets_and_data():
    summary = Summary(total=1, by_source={"ServiceDesk": 1}, by_status={"DONE": 1},
                      auto=1, manual=0, unaddressed_total=0,
                      unaddressed_by_source={}, cc_yes=1, cc_no=0)
    data = build_report_xlsx([_row()], summary)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Сводка", "Сообщения"]
    detail = wb["Сообщения"]
    assert detail["A1"].value == "ID письма"
    assert detail["F2"].value == "DONE"
    summary_sheet = wb["Сводка"]
    flat = [c.value for col in summary_sheet.iter_cols() for c in col]
    assert "Всего сообщений:" in flat
    assert 1 in flat
